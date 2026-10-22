"""Every formatting feature of the brief, driven through the real toolbar, menus and shortcuts in a browser, then
saved, reloaded and exported."""
from __future__ import annotations

import io
import re

import pytest
from docx import Document as WordDocument

from .browser import app_server
from .test_e2e import EDITOR, html, open_app, wait_html

pytestmark = pytest.mark.browser


def btn(page, name: str):
    return page.get_by_role("button", name=re.compile(rf"^{re.escape(name)}"))


def new_doc(page, title: str = "Feature test") -> None:
    page.get_by_role("button", name="New document").click()
    page.locator(".q-dialog input").first.fill(title)
    page.get_by_role("button", name="Create").click()
    page.wait_for_selector(f".ws-title:has-text('{title}')")
    page.wait_for_function("document.querySelector('.cx-root').__cxController.getHtml() === '<p></p>'")


def fresh_paragraph(page, text: str, select: bool = True) -> None:
    """Replace the document with one paragraph; the text is selected unless select=False (caret at the end)."""
    page.click(EDITOR)
    page.keyboard.press("ControlOrMeta+a")
    page.keyboard.press("Backspace")
    page.keyboard.type(text)
    if select:
        page.keyboard.press("ControlOrMeta+a")


@pytest.fixture
def owner_page(chromium):
    with app_server() as (url, db):
        page = open_app(chromium, url)
        new_doc(page)
        yield page, url
        page.context.close()


MARKS = [("Bold", "<strong>"), ("Italic", "<em>"), ("Underline", "<u>"), ("Strikethrough", "<s>"),
         ("Subscript", "<sub>"), ("Superscript", "<sup>"), ("Inline code", "<code>")]


@pytest.mark.parametrize("label,tag", MARKS)
def test_text_marks_toggle_with_toolbar(owner_page, label, tag):
    page, _ = owner_page
    fresh_paragraph(page, "sample")
    btn(page, label).click()
    assert tag in html(page)
    btn(page, label).click()
    assert tag not in html(page)


@pytest.mark.parametrize("keys,tag", [("ControlOrMeta+b", "<strong>"), ("ControlOrMeta+i", "<em>"), ("ControlOrMeta+u", "<u>")])
def test_word_style_shortcuts(owner_page, keys, tag):
    page, _ = owner_page
    fresh_paragraph(page, "sample")
    page.keyboard.press(keys)
    assert tag in html(page)


def test_color_and_highlight(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "sample")
    btn(page, "Text color").click()
    page.locator(".cx-swatches .cx-swatch").nth(1).click()
    assert re.search(r'<span style="color: rgb\(\d+, \d+, \d+\);">sample</span>', html(page))
    fresh_paragraph(page, "other")
    btn(page, "Highlight").click()
    page.locator(".cx-swatches .cx-swatch").nth(0).click()
    assert re.search(r'<mark style="background-color: rgb\(\d+, \d+, \d+\);">', html(page))


def test_headings_h1_to_h4_and_paragraph(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "Title")
    for n in (1, 2, 3, 4):
        page.select_option("select[aria-label='Text style']", f"h{n}")
        assert f"<h{n}>Title</h{n}>" in html(page)
    page.select_option("select[aria-label='Text style']", "p")
    assert "<p>Title</p>" in html(page)


def test_alignment(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "aligned")
    for value in ("center", "right", "justify"):
        btn(page, f"Align {value}").click()
        assert f'style="text-align: {value};"' in html(page)


def test_lists_checklist_indent_outdent(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "one", select=False)
    btn(page, "Bulleted list").click()
    page.keyboard.press("Enter")
    page.keyboard.type("two")
    btn(page, "Indent").click()
    assert re.search(r"<ul><li><p>one</p><ul><li><p>two</p></li></ul></li></ul>", html(page))
    btn(page, "Outdent").click()
    assert "<li><p>one</p></li><li><p>two</p></li>" in html(page)
    btn(page, "Numbered list").click()
    assert "<ol>" in html(page)
    btn(page, "Checklist").click()
    assert 'data-type="taskList"' in html(page) and 'data-type="taskItem"' in html(page)
    page.locator(".cx-canvas li[data-type='taskItem'] input").first.check()
    assert 'data-checked="true"' in html(page)


def test_quote_code_block_divider(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "quoted", select=False)
    btn(page, "Quote").click()
    assert "<blockquote>" in html(page)
    btn(page, "Quote").click()
    btn(page, "Code block").click()
    assert "<pre><code>" in html(page)
    btn(page, "Code block").click()
    btn(page, "Divider").click()
    assert "<hr>" in html(page)


def test_link_add_edit_remove_and_open_in_new_tab(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "site")
    btn(page, "Link").click()
    page.locator(".cx-modal input").fill("example.com")
    page.get_by_role("button", name="Apply").click()
    assert '<a href="https://example.com" target="_blank" rel="noopener noreferrer">site</a>' in html(page)
    page.click(f"{EDITOR} a")
    page.keyboard.press("ControlOrMeta+k")
    page.locator(".cx-modal input").fill("https://example.org/page")
    page.get_by_role("button", name="Apply").click()
    assert 'href="https://example.org/page"' in html(page)
    page.click(f"{EDITOR} a")
    page.keyboard.press("ControlOrMeta+k")
    page.get_by_role("button", name="Remove link").click()
    assert "<a " not in html(page)
    page.keyboard.press("ControlOrMeta+a")
    btn(page, "Link").click()
    page.locator(".cx-modal input").fill("javascript:alert(1)")
    page.get_by_role("button", name="Apply").click()
    assert "Only http, https and mailto" in page.locator(".cx-modal").inner_text()


def test_table_rows_columns_header_merge(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "")
    btn(page, "Insert table").click()
    assert "<table>" in html(page) and html(page).count("<th>") == 3
    page.locator(f"{EDITOR} td").first.click()
    btn(page, "Add row below").click()
    btn(page, "Add column after").click()
    assert html(page).count("<tr>") == 4 and html(page).count("<th>") == 4
    btn(page, "Toggle header row").click()
    assert "<th>" not in html(page)
    btn(page, "Toggle header row").click()
    cells = page.locator(f"{EDITOR} td")
    cells.nth(0).click()
    cells.nth(1).click(modifiers=["Shift"])
    btn(page, "Merge cells").click()
    assert "colspan=\"2\"" in html(page)
    btn(page, "Split cell").click()
    assert "colspan" not in html(page)
    btn(page, "Delete row").click()
    btn(page, "Delete column").click()
    btn(page, "Delete table").click()
    assert "<table>" not in html(page)


def test_undo_redo_and_clear_formatting(owner_page):
    page, _ = owner_page
    fresh_paragraph(page, "sample")
    page.keyboard.press("ControlOrMeta+b")
    page.keyboard.press("ControlOrMeta+i")
    btn(page, "Clear formatting").click()
    assert "<strong>" not in html(page) and "<em>" not in html(page)
    btn(page, "Undo").click()
    assert "<strong>" in html(page)
    btn(page, "Redo").click()
    assert "<strong>" not in html(page)


def test_slash_menu_inserts_blocks(owner_page):
    page, _ = owner_page
    page.click(EDITOR)
    page.keyboard.type("/heading 2")
    page.wait_for_selector(".cx-slash .cx-slash-item")
    page.keyboard.press("Enter")
    page.keyboard.type("Section")
    assert "<h2>Section</h2>" in html(page)
    page.keyboard.press("Enter")
    page.keyboard.type("/table")
    page.keyboard.press("Enter")
    assert "<table>" in html(page)


def test_formatting_survives_save_reload_and_word_export(owner_page):
    page, url = owner_page
    page.click(EDITOR)
    page.keyboard.type("/heading 1")
    page.keyboard.press("Enter")
    page.keyboard.type("Report")
    page.keyboard.press("Enter")
    page.keyboard.type("Plain then ")
    page.keyboard.press("ControlOrMeta+b")
    page.keyboard.type("bold")
    page.keyboard.press("ControlOrMeta+b")
    page.keyboard.press("Enter")
    page.keyboard.type("/bulleted")
    page.keyboard.press("Enter")
    page.keyboard.type("first")
    page.keyboard.press("Enter")
    page.keyboard.type("second")
    page.keyboard.press("Enter")
    page.keyboard.press("Enter")
    page.keyboard.type("/table")
    page.keyboard.press("Enter")
    before = html(page)
    page.get_by_role("button", name=re.compile("^Save as new version")).click()
    page.wait_for_selector("text=Saved as version 2")
    page.reload()
    page.wait_for_selector(".ws-doc-title:has-text('Feature test')")
    page.locator(".ws-doc-title:has-text('Feature test')").click()  # reload opens the first document
    wait_html(page, "Report")
    assert html(page) == before
    page.get_by_role("button", name="Export").click()
    with page.expect_download() as info:
        page.get_by_text("Word (.docx)").click()
    doc = WordDocument(io.BytesIO(open(info.value.path(), "rb").read()))
    styles = [p.style.name for p in doc.paragraphs]
    assert "Heading 1" in styles and any("List" in s for s in styles) and len(doc.tables) >= 2
    assert any(any(r.bold for r in p.runs if r.text == "bold") for p in doc.paragraphs)
