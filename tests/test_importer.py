"""File import: every supported type becomes sanitized, editable HTML; hostile files are refused."""
from __future__ import annotations

import io
import struct
import zipfile
import zlib

import pytest
from docx import Document as WordDocument
from fpdf import FPDF
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from ctxeditor.errors import EditorError
from ctxeditor.importer import MAX_BYTES, UnsupportedFile, convert_safe, convert_upload
from ctxeditor.sanitizer import sanitize_html

def _png() -> bytes:
    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b""))


PNG = _png()


def make_docx() -> bytes:
    doc = WordDocument()
    doc.add_heading("Project brief", level=1)
    doc.add_heading("Scope", level=2)
    doc.add_heading("Deep heading", level=5)
    p = doc.add_paragraph("Plain, ")
    p.add_run("bold").bold = True
    p.add_run(" and ")
    p.add_run("italic").italic = True
    p.add_run(" and ")
    p.add_run("underlined").underline = True
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Second point", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Risk", "Impact"
    table.cell(1, 0).text, table.cell(1, 1).text = "Lock-in", "High"
    doc.add_picture(io.BytesIO(PNG))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Cost"])
    ws.append(["Licences", 1200.0])
    ws.append(["Hosting", 99.5])
    hidden = wb.create_sheet("Secret")
    hidden.append(["do not import"])
    hidden.sheet_state = "hidden"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def make_pptx() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Roadmap"
    frame = slide.placeholders[1].text_frame
    frame.text = "Q1 discovery"
    sub = frame.add_paragraph()
    sub.text = "interviews"
    sub.level = 1
    frame.add_paragraph().text = "Q2 build"
    slide.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(5))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def make_pdf(scanned: bool = False) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    if not scanned:
        pdf.set_font("Helvetica", "B", 24)
        pdf.cell(0, 14, "Annual Report", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, "This is the first paragraph of the report. It continues over several lines so that the "
                             "importer has to join them back into one paragraph again.")
        pdf.ln(4)
        pdf.multi_cell(0, 6, "Second paragraph here.")
        pdf.ln(4)
        pdf.cell(0, 6, "- first bullet", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 6, "- second bullet", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def doc_of(imported) -> str:
    assert sanitize_html(imported.html) == imported.html, "importer output must already be sanitized"
    return imported.html


def test_docx_structure_formatting_lists_tables_and_images():
    imported = convert_upload("Project_brief.docx", make_docx())
    out = doc_of(imported)
    assert imported.title == "Project brief" and imported.kind == "docx"
    assert "<h1>Project brief</h1>" in out and "<h2>Scope</h2>" in out
    assert "<h4>Deep heading</h4>" in out  # h5 is mapped onto the editor's deepest level
    assert "<strong>bold</strong>" in out and "<em>italic</em>" in out and "<u>underlined</u>" in out
    assert "<ul><li>First point</li><li>Second point</li></ul>" in out and "<ol><li>Step one</li></ol>" in out
    assert "<table>" in out and "Lock-in" in out
    assert "<img" not in out
    assert any("1 image was not imported" in n for n in imported.notes)
    assert imported.tables == 1 and imported.headings >= 3 and imported.words > 10


def test_xlsx_sheets_become_headings_and_tables_hidden_sheet_skipped():
    imported = convert_upload("budget.xlsx", make_xlsx())
    out = doc_of(imported)
    assert "<h2>Budget</h2>" in out and "<th><p>Item</p></th>" in out
    assert "<td><p>1200</p></td>" in out and "<td><p>99.5</p></td>" in out
    assert "do not import" not in out and "Secret" not in out
    assert any("Hidden sheets" in n for n in imported.notes)


def test_pptx_slides_become_sections_with_nested_bullets():
    imported = convert_upload("deck.pptx", make_pptx())
    out = doc_of(imported)
    assert "<h2>Roadmap</h2>" in out
    assert "<li><p>Q1 discovery</p><ul><li><p>interviews</p></li></ul></li>" in out
    assert "<li><p>Q2 build</p></li>" in out
    assert any("picture" in n for n in imported.notes)


def test_pdf_headings_paragraphs_and_bullets():
    imported = convert_upload("report.pdf", make_pdf())
    out = doc_of(imported)
    assert "<h1>Annual Report</h1>" in out
    assert "first paragraph of the report. It continues over several lines" in out  # lines joined into one paragraph
    assert "<p>Second paragraph here.</p>" in out
    assert "<ul><li><p>first bullet</p></li><li><p>second bullet</p></li></ul>" in out


def test_scanned_pdf_is_refused_with_a_clear_message():
    with pytest.raises(UnsupportedFile, match="no text layer"):
        convert_upload("scan.pdf", make_pdf(scanned=True))


def test_markdown_tables_code_checklist_mermaid_and_raw_html_escaped():
    md = (
        "# Title\n\nSome *emphasis*, **strong**, ~~gone~~ and `code`.\n\n"
        "| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "- [x] done thing\n- [ ] open thing\n\n"
        "```python\nprint('hi')\n```\n\n"
        "```mermaid\ngraph TD; A-->B\n```\n\n"
        "<script>alert(1)</script>\n\n<b onclick='x()'>raw</b>\n\n![pic](http://x/y.png)\n"
    ).encode()
    imported = convert_upload("notes.md", md)
    out = doc_of(imported)
    assert "<h1>Title</h1>" in out and "<em>emphasis</em>" in out and "<s>gone</s>" in out
    assert "<table>" in out and "<th>A</th>" in out.replace("<p>", "").replace("</p>", "")
    assert '<ul data-type="taskList">' in out and 'data-type="taskItem" data-checked="true"' in out
    assert "done thing" in out and "[x]" not in out and "[ ]" not in out
    assert '<pre><code class="language-python">' in out
    assert '<div data-type="mermaid">graph TD; A--&gt;B</div>' in out
    from lxml import html as lh
    tree = lh.fragment_fromstring(out, create_parent="div")
    assert not tree.xpath(".//script") and not [e for e in tree.iter() if "onclick" in getattr(e, "attrib", {})]
    assert "alert(1)" in tree.text_content()  # shown as plain text, never run
    assert "<img" not in out


def test_html_file_is_sanitized_and_scripts_dropped():
    page = (b"<html><head><title>x</title><style>p{color:red}</style><script>evil()</script></head><body>"
            b"<h1>Hello</h1><p onclick='evil()'>Text <a href='javascript:evil()'>bad</a> "
            b"<a href='https://ok.example'>ok</a></p><iframe src='https://evil'></iframe></body></html>")
    out = doc_of(convert_upload("page.html", page))
    assert "<h1>Hello</h1>" in out and "evil" not in out and "iframe" not in out and "javascript" not in out
    assert 'href="https://ok.example"' in out


def test_csv_with_detected_delimiter_and_truncation_note():
    rows = "\n".join(f"{i};name{i};{i * 2}" for i in range(600))
    imported = convert_upload("data.csv", ("id;name;value\n" + rows).encode())
    out = doc_of(imported)
    assert out.count("<tr>") == 500 and "<th><p>id</p></th>" in out
    assert any("cut to 500 rows" in n for n in imported.notes)


def test_plain_text_paragraphs_and_line_breaks_with_escaping():
    out = doc_of(convert_upload("readme.txt", b"First line\nsecond line\n\nNew <b>paragraph</b> & more\n"))
    assert out == "<p>First line<br>second line</p><p>New &lt;b&gt;paragraph&lt;/b&gt; &amp; more</p>"


def test_text_encodings():
    assert "café" in convert_upload("a.txt", "café".encode("cp1252")).html
    assert "café" in convert_upload("a.txt", "\ufeffcafé".encode("utf-8")).html


@pytest.mark.parametrize("name,data,message", [
    ("old.doc", b"x", "old format"),
    ("old.xls", b"x", "old format"),
    ("tool.exe", b"MZ", "Upload a Word"),
    ("noext", b"hello", "Upload a Word"),
    ("empty.txt", b"", "empty"),
    ("fake.docx", b"just some text", "not a real .docx"),
    ("fake.pdf", b"not a pdf", "not a PDF"),
    ("binary.txt", b"abc\x00\x01\x02def", "binary"),
    ("huge.txt", b"a" * (MAX_BYTES + 1), "larger than"),
    ("blank.csv", b"\n\n", "No content"),
])
def test_refusals(name, data, message):
    with pytest.raises(UnsupportedFile, match=message):
        convert_upload(name, data)


def test_zip_that_is_not_a_docx_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("hello.txt", "hi")
    with pytest.raises(UnsupportedFile, match="not a real .docx"):
        convert_upload("a.docx", out.getvalue())


def test_decompression_bomb_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w:document xmlns:w='x'/>")
        z.writestr("word/bomb.bin", b"\x00" * (60 * 1024 * 1024))
    with pytest.raises(UnsupportedFile, match="bomb"):
        convert_upload("bomb.docx", out.getvalue())


def test_xml_entity_attack_in_docx_is_refused():
    out = io.BytesIO()
    evil = ('<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "boom"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;">]>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
            '<w:p><w:r><w:t>&b;</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", evil)
    with pytest.raises(UnsupportedFile, match="unsafe XML"):
        convert_upload("evil.docx", out.getvalue())


def test_xxe_external_entity_is_refused():
    out = io.BytesIO()
    evil = ('<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
            '<w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", evil)
    with pytest.raises(UnsupportedFile, match="unsafe XML"):
        convert_upload("xxe.docx", out.getvalue())


def test_convert_safe_returns_messages_not_exceptions():
    imported, error = convert_safe("a.docx", b"junk")
    assert imported is None and "not a real .docx" in error
    imported, error = convert_safe("ok.txt", b"fine")
    assert error is None and imported.html == "<p>fine</p>"
    # a damaged but well-formed archive must not leak a traceback
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("word/document.xml", "<not-closed>")
    imported, error = convert_safe("bad.docx", out.getvalue())
    assert imported is None and error and "Traceback" not in error


def test_file_names_are_reduced_to_a_safe_title():
    assert convert_upload("../../etc/Q3_plan\x00.txt", b"x").title == "Q3 plan"
    assert convert_upload("C:\\temp\\my_file.txt", b"x").title == "my file"
    assert len(convert_upload(("a" * 400) + ".txt", b"x").title) <= 200


def test_oversized_conversion_is_refused_not_truncated():
    with pytest.raises(UnsupportedFile, match="too large"):
        convert_upload("big.txt", ("word " * 400_000).encode()[:MAX_BYTES])


def test_imported_html_can_be_stored_as_a_document(env):
    from ctxeditor.errors import PermissionDenied
    imported = convert_upload("brief.docx", make_docx())
    created = env.docs.create_document(env.actor("priya"), env.seeded.projects["acme"], imported.title, "Imported",
                                       imported.html)
    assert created.version_number == 1 and created.is_draft_view and created.editable
    assert "<h1>Project brief</h1>" in created.html
    with pytest.raises(PermissionDenied):
        env.docs.create_document(env.actor("sana"), env.seeded.projects["acme"], imported.title, "Imported",
                                 imported.html)
    assert env.docs.can_create(env.actor("priya"), env.seeded.projects["acme"]) is True
    assert env.docs.can_create(env.actor("sana"), env.seeded.projects["acme"]) is False
