"""DOCX builder: structure checks on the reopened file. No browser needed."""
from __future__ import annotations

import io
import re
import struct
import zlib

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, RGBColor

from ctxeditor.export.docx import build_docx
from ctxeditor.models import DocMetadata, ExportPayload
from ctxeditor.sample_content import SAMPLE_HTML

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def png(width: int = 4, height: int = 3) -> bytes:
    """A valid tiny RGB PNG, built by hand (no Pillow)."""
    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
    rows = b"".join(b"\x00" + b"\x20\x80\xc0" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def payload(html: str = SAMPLE_HTML, *, draft: bool = False, published_at: str | None = None) -> ExportPayload:
    meta = DocMetadata(
        title="Showcase", doc_type="Guide", version_number=3, status="Draft" if draft else "Published",
        created_at="2026-10-01T10:00:00.000000Z", created_by="Priya Nair",
        last_edited_at="2026-10-02T11:30:00.000000Z", last_edited_by="Arjun Mehta",
        last_published_at=published_at, last_published_by="Priya Nair" if published_at else None)
    return ExportPayload(html=html, metadata=meta, is_draft=draft, title=meta.title)


def build(html: str = SAMPLE_HTML, images: dict[int, bytes] | None = None, **kwargs):
    tz = kwargs.pop("tz", None)
    data = build_docx(payload(html, **kwargs), images or {}, tz)
    return Document(io.BytesIO(data))


def all_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in doc.sections:
        parts.extend(p.text for p in section.header.paragraphs)
        parts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(parts)


def runs_with(doc, text: str):
    return [r for r in doc.element.body.iter(qn("w:r")) if "".join(t.text or "" for t in r.iter(qn("w:t"))) == text]


@pytest.fixture(scope="module")
def doc():
    return build(images={0: png(), 1: png(8, 6)})


def test_headings_use_real_word_styles(doc):
    found = {(p.style.name, p.text) for p in doc.paragraphs}
    assert ("Heading 1", "Customer Portal: context editor showcase") in found
    assert ("Heading 2", "Goals") in found
    assert ("Heading 3", "Risks") in found
    assert ("Heading 4", "Sign-in flow") in found
    assert ("Title", "Showcase") in found


def test_page_is_a4_with_margins(doc):
    section = doc.sections[0]
    assert abs(section.page_width - Cm(21)) < 1000 and abs(section.page_height - Cm(29.7)) < 1000
    assert abs(section.left_margin - Cm(2.5)) < 1000


def test_body_font_is_cambria(doc):
    assert doc.styles["Normal"].font.name == "Cambria"
    assert doc.styles["Normal"].font.size.pt == 11


def test_list_styles_and_real_numbering(doc):
    by_text = {p.text: p for p in doc.paragraphs}
    assert by_text["Single place for project context"].style.name == "List Bullet"
    assert by_text["Draft and publish"].style.name == "List Bullet 2"
    assert by_text["Write"].style.name == "List Number"
    for text in ("Single place for project context", "Draft and publish", "Write", "Publish"):
        assert by_text[text]._p.pPr.numPr is not None
    nums = {by_text[t]._p.pPr.numPr.numId.val for t in ("Write", "Review", "Publish")}
    assert len(nums) == 1  # one ordered list shares one numbering instance


def test_ordered_lists_restart_numbering():
    html = "<ol><li><p>a</p></li></ol><p>between</p><ol><li><p>b</p></li></ol>"
    d = build(html)
    numbered = [p for p in d.paragraphs if p.style.name == "List Number"]
    assert len(numbered) == 2
    assert numbered[0]._p.pPr.numPr.numId.val != numbered[1]._p.pPr.numPr.numId.val
    xml = d.part.numbering_part.element.xml
    assert xml.count("w:startOverride") >= 2


def test_ordered_list_start_attribute():
    d = build('<ol start="5"><li><p>five</p></li></ol>')
    assert 'w:startOverride w:val="5"' in d.part.numbering_part.element.xml


def test_task_list_uses_ballot_boxes(doc):
    texts = [p.text for p in doc.paragraphs]
    assert "☑ Draft the PRD" in texts
    assert "☐ Collect feedback" in texts


def risks_table(d):
    return next(t for t in d.tables if "Vendor lock-in" in t._tbl.xml)


def test_table_header_row_repeats_and_is_shaded(doc):
    table = risks_table(doc)
    first_row = table.rows[0]._tr
    assert first_row.find(f"{W}trPr/{W}tblHeader") is not None
    assert table.rows[1]._tr.find(f"{W}trPr") is None
    assert "F2F2F2" in table.rows[0].cells[0]._tc.xml
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold


def test_table_colspan_becomes_merged_cell(doc):
    table = risks_table(doc)
    assert re.search(r'<w:gridSpan w:val="2"', table._tbl.xml)
    assert table.rows[2].cells[0].text == "Merged cell"
    assert len(table.columns) == 3


def test_table_rowspan_becomes_vertical_merge():
    html = ("<table><tbody><tr><td rowspan=\"2\"><p>tall</p></td><td><p>a</p></td></tr>"
            "<tr><td><p>b</p></td></tr></tbody></table>")
    d = build(html)
    xml = d.tables[-1]._tbl.xml
    assert "w:vMerge" in xml and "tall" in xml


def test_malformed_overlapping_spans_do_not_crash():
    html = ("<table><tr><td rowspan=\"2\"><p>a</p></td><td colspan=\"5\"><p>b</p></td></tr>"
            "<tr><td colspan=\"3\"><p>c</p></td></tr></table>")
    assert build(html).tables


def test_hyperlink_is_relationship_backed(doc):
    assert doc.element.body.find(f".//{W}hyperlink") is not None
    links = [r for r in doc.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert [(r.target_ref, r.is_external) for r in links] == [("https://example.com", True)]
    link_run = doc.element.body.find(f".//{W}hyperlink/{W}r")
    assert link_run is not None and "link" in "".join(t.text for t in link_run.iter(qn("w:t")))


def test_inline_marks_become_runs(doc):
    def run(text):
        return next(r for p in doc.paragraphs for r in p.runs if r.text == text)
    assert run("bold").bold and run("italic").italic and run("underline").underline
    assert run("strikethrough").font.strike
    sub = [r for p in doc.paragraphs for r in p.runs if r.text == "2" and r.font.subscript]
    sup = [r for p in doc.paragraphs for r in p.runs if r.text == "2" and r.font.superscript]
    assert sub and sup
    assert run("inline code").font.name == "Consolas"


def test_color_and_highlight(doc):
    colored = next(r for p in doc.paragraphs for r in p.runs if r.text == "colored text")
    assert colored.font.color.rgb == RGBColor(0xC6, 0x28, 0x28)
    highlighted = next(r for p in doc.paragraphs for r in p.runs if r.text == "highlighted text")
    assert 'w:fill="FFF59D"' in highlighted._r.xml


def test_rgb_colour_notation():
    d = build('<p><span style="color: rgb(0, 128, 255)">x</span></p>')
    run = next(r for p in d.paragraphs for r in p.runs if r.text == "x")
    assert run.font.color.rgb == RGBColor(0, 128, 255)


def test_paragraph_alignment(doc):
    para = next(p for p in doc.paragraphs if p.text == "A centered paragraph.")
    assert para.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_blockquote_and_rule(doc):
    quote = next(p for p in doc.paragraphs if p.text == "Good context makes good decisions.")
    assert quote.style.name == "Quote"
    rules = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:pBdr")) is not None]
    assert rules and rules[0].text == ""


def test_code_block_is_monospace_style_with_colours(doc):
    code = next(p for p in doc.paragraphs if "def greet(name):" in p.text)
    assert code.style.name == "Code Block" and code.style.font.name == "Consolas"
    assert any(r.font.color is not None and r.font.color.rgb is not None for r in code.runs)
    assert "return" in code.text and code._p.find(f".//{W}br") is not None


def test_code_block_without_language_is_plain():
    d = build("<pre><code>a = 1\nb = 2</code></pre>")
    code = next(p for p in d.paragraphs if p.style.name == "Code Block")
    assert "a = 1" in code.text and "b = 2" in code.text


def test_two_inline_pictures_with_alt_text(doc):
    assert len(doc.inline_shapes) == 2
    descrs = [shape._inline.docPr.get("descr") for shape in doc.inline_shapes]
    assert "sequenceDiagram" in descrs[0] and "U->>A: Open the portal" in descrs[0]
    assert '"type":"bar"' in descrs[1]
    assert all(shape.width <= Cm(16) + 10 for shape in doc.inline_shapes)


def test_picture_ordinals_count_both_kinds_in_document_order():
    d = build(SAMPLE_HTML, {1: png()})  # only the chart (second block) has an image
    assert len(d.inline_shapes) == 1
    assert '"type":"bar"' in d.inline_shapes[0]._inline.docPr.get("descr")
    assert "sequenceDiagram" in all_text(d)  # the diagram fell back to its source


def test_missing_images_become_placeholders():
    d = build()
    assert len(d.inline_shapes) == 0
    text = all_text(d)
    assert "[Diagram image unavailable]" in text and "sequenceDiagram" in text
    assert "[Chart image unavailable: Quarterly signups]" in text


def test_unreadable_image_becomes_placeholder():
    d = build(images={0: b"not a png", 1: png()})
    assert len(d.inline_shapes) == 1 and "[Diagram image unavailable]" in all_text(d)


def test_footer_has_page_x_of_y_fields(doc):
    xml = doc.sections[0].footer._element.xml
    assert "PAGE" in xml and "NUMPAGES" in xml
    assert doc.sections[0].footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "Page" in doc.sections[0].footer.paragraphs[0].text


def test_header_carries_title(doc):
    assert doc.sections[0].header.paragraphs[0].text.startswith("Showcase")


def test_draft_marker_only_for_drafts():
    draft = build(draft=True)
    assert "DRAFT" in "\n".join(p.text for p in draft.paragraphs)
    assert "DRAFT" in draft.sections[0].header.paragraphs[0].text
    final = build()
    assert "DRAFT" not in all_text(final)


def test_metadata_table_content():
    d = build(published_at="2026-10-03T09:00:00.000000Z", tz="UTC")
    meta = {row.cells[0].text: row.cells[1].text for row in d.tables[0].rows}
    assert meta["Title"] == "Showcase" and meta["Status"] == "Published" and meta["Type"] == "Guide"
    assert meta["Version"] == "3" and meta["Created by"] == "Priya Nair" and meta["Last edited by"] == "Arjun Mehta"
    assert meta["Created"] == "Oct 1, 2026, 10:00 AM"
    assert meta["Last published"] == "Oct 3, 2026, 9:00 AM" and meta["Published by"] == "Priya Nair"


def test_metadata_unpublished_and_viewer_timezone():
    d = build(draft=True, tz="Asia/Kolkata")
    meta = {row.cells[0].text: row.cells[1].text for row in d.tables[0].rows}
    assert meta["Status"] == "Draft" and meta["Last published"] == "-" and meta["Published by"] == "-"
    assert meta["Created"] == "Oct 1, 2026, 3:30 PM"


def test_core_properties(doc):
    props = doc.core_properties
    assert props.title == "Showcase" and props.author == "Arjun Mehta" and props.subject == "Guide"


def test_empty_html_gives_only_the_metadata_block():
    for html in ("", "   ", "<p></p>"):
        d = build(html)
        assert len(d.tables) == 1
        assert {row.cells[0].text for row in d.tables[0].rows} >= {"Title", "Status"}


def test_unknown_and_stray_markup_is_walked_as_container():
    d = build("loose text <custom><p>inside custom</p></custom><div>in a div <strong>bold</strong></div>")
    text = all_text(d)
    assert "loose text" in text and "inside custom" in text and "in a div" in text


def test_whitespace_collapses_like_html():
    d = build("<p>  a \n   b   <strong> c </strong>  d  </p>")
    assert next(p.text for p in d.paragraphs if p.text.startswith("a")) == "a b c d"


def test_br_becomes_line_break():
    d = build("<p>one<br>two</p>")
    para = next(p for p in d.paragraphs if p.text.startswith("one"))
    assert para._p.find(f".//{W}br") is not None


def test_nested_content_in_table_cell_and_list_item():
    html = ("<table><tr><td><ul><li><p>in cell list</p></li></ul><p>after</p></td></tr></table>"
            "<ul><li><p>item</p><p>continued</p><pre><code>x = 1</code></pre></li></ul>")
    d = build(html)
    cell = d.tables[-1].rows[0].cells[0]
    assert [p.style.name for p in cell.paragraphs][:1] == ["List Bullet"]
    assert any(p.text == "after" for p in cell.paragraphs)
    assert next(p for p in d.paragraphs if p.text == "continued").style.name != "List Bullet"


def test_document_ends_with_a_paragraph_after_a_final_table():
    d = build("<table><tr><td><p>x</p></td></tr></table>")
    body = [c for c in d.element.body if c.tag != qn("w:sectPr")]
    assert body[-1].tag == qn("w:p")


def test_unsafe_link_scheme_is_not_linked():
    d = build('<p><a href="javascript:alert(1)">bad</a></p>')
    assert not [r for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "bad" in all_text(d)


def test_numbering_part_is_schema_ordered(doc):
    root = doc.part.numbering_part.element
    tags = [c.tag for c in root]
    last_abstract = max(i for i, t in enumerate(tags) if t == qn("w:abstractNum"))
    first_num = min(i for i, t in enumerate(tags) if t == qn("w:num"))
    assert last_abstract < first_num
