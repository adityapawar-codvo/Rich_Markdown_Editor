"""PDF and Word export through headless Chromium. Permission rules are covered in test_documents.py; here we prove
the pipeline renders formatting, tables, links, diagrams and charts, marks drafts, and refuses what it must."""
from __future__ import annotations

import io

import pytest
from docx import Document as WordDocument
from pypdf import PdfReader

from ctxeditor.errors import PermissionDenied
from ctxeditor.export.browser import BrowserPool, ExportFailed
from ctxeditor.export.service import ExportService

pytestmark = pytest.mark.browser


@pytest.fixture
async def exporter(env):
    pool = BrowserPool()
    try:
        await pool._ensure()
    except ExportFailed:
        pytest.skip("Chromium is not installed (run: uv run playwright install chromium)")
    yield ExportService(env.docs, pool)
    await pool.close()


async def test_pdf_of_published_version_has_metadata_content_and_page_numbers(env, exporter):
    result = await exporter.export(env.actor("sana"), env.doc_id, "pdf", tz="Asia/Kolkata")
    assert result.filename == "context-editor-showcase-v1.pdf" and result.media_type == "application/pdf"
    reader = PdfReader(io.BytesIO(result.data))
    text = "\n".join(page.extract_text() for page in reader.pages)
    for needle in ("Context editor showcase", "Published", "Priya Nair", "Goals", "Vendor lock-in", "Merged cell",
                   "Open the portal", "Page 1 of"):
        assert needle in text, needle
    assert "DRAFT" not in text
    assert "Open questions" not in text  # the draft's content never reaches a team member's export
    assert sum(len(page.images) for page in reader.pages) >= 1  # the chart is drawn on a canvas


async def test_owner_draft_export_is_marked_draft(env, exporter):
    pdf = await exporter.export(env.actor("priya"), env.doc_id, "pdf")
    assert pdf.filename.endswith("-v2-DRAFT.pdf")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    assert "DRAFT" in text and "Open questions" in text
    docx = await exporter.export(env.actor("priya"), env.doc_id, "docx")
    assert docx.filename.endswith("-v2-DRAFT.docx")
    assert any("DRAFT" in p.text for p in WordDocument(io.BytesIO(docx.data)).paragraphs)


async def test_docx_has_real_structure_and_diagram_and_chart_pictures(env, exporter):
    result = await exporter.export(env.actor("dmitri"), env.doc_id, "docx")
    doc = WordDocument(io.BytesIO(result.data))
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Heading 1", "Heading 2", "Heading 3", "Heading 4"} <= styles
    assert len(doc.tables) >= 2  # metadata block + the risks table
    assert len(doc.inline_shapes) == 2  # the Mermaid diagram and the chart, rendered by the same bundle


async def test_team_cannot_export_draft(env, exporter):
    with pytest.raises(PermissionDenied):
        await exporter.export(env.actor("sana"), env.doc_id, "pdf", version=2)


async def test_exports_are_audited(env, exporter):
    await exporter.export(env.actor("erin"), env.doc_id, "pdf")
    with env.db.read() as conn:
        rows = conn.execute("select metadata from audit_events where action = 'doc_exported'").fetchall()
    assert len(rows) == 1 and '"format": "pdf"' in rows[0]["metadata"]


async def test_bad_mermaid_does_not_break_export(env, exporter):
    owner = env.actor("priya")
    env.docs.save_draft(owner, env.doc_id, 2, '<p>before</p><div data-type="mermaid">this is not mermaid ((</div><p>after</p>')
    pdf = await exporter.export(owner, env.doc_id, "pdf")
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.data)).pages)
    assert "before" in text and "after" in text
