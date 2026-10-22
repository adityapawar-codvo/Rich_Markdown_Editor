"""Turn an uploaded file into editor HTML so it opens in the editor and stays fully editable.

Supported: Word (.docx), PowerPoint (.pptx), Excel (.xlsx), PDF with a text layer, Markdown, HTML, CSV and plain text.
Images are not imported (the editor has no images); the user is told how many were skipped.

The upload is treated as hostile input:
- the type is decided by extension and then checked against the file's signature;
- Office files are opened through an archive guard (entry count, unzipped size) and every XML part is parsed with
  defusedxml first, so entity and DTD tricks are refused;
- the output of every converter goes through the same sanitizer as everything else the editor stores, so a file can
  never put script, styles or unknown markup into a document.

convert_upload() is pure (no I/O, no UI), so it can run in a worker process; convert_safe() reports a failure as a
value because an exception from a worker process would lose its message.
"""
from __future__ import annotations

import csv
import html as html_lib
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass
from pathlib import PurePath

from lxml import html as lxml_html

from .errors import EditorError, ValidationFailed
from .sanitizer import sanitize_html

MAX_BYTES = 10 * 1024 * 1024
MAX_ZIP_ENTRIES = 3000
MAX_UNZIPPED = 150 * 1024 * 1024
MAX_XML_PART = 40 * 1024 * 1024
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 50
MAX_SHEETS = 20
MAX_SLIDES = 300
MAX_PDF_PAGES = 200
MAX_TITLE = 200

SUPPORTED = (".docx", ".pptx", ".xlsx", ".pdf", ".md", ".markdown", ".html", ".htm", ".csv", ".txt")
ACCEPT = ",".join(SUPPORTED)
LEGACY = {".doc": ".docx", ".ppt": ".pptx", ".xls": ".xlsx", ".rtf": ".docx", ".odt": ".docx"}


class UnsupportedFile(EditorError):
    message = "This kind of file cannot be imported."


@dataclass(frozen=True)
class Imported:
    title: str
    html: str  # sanitized, ready for DocumentService.create_document
    notes: tuple[str, ...]  # things the user should know (images skipped, rows cut off, ...)
    kind: str
    words: int
    headings: int
    tables: int


# ---------------------------------------------------------------------------------------------------- entry points

def convert_safe(filename: str, data: bytes) -> tuple[Imported | None, str | None]:
    """For worker processes: (result, None) or (None, message safe to show the user)."""
    try:
        return convert_upload(filename, data), None
    except EditorError as exc:
        return None, exc.message
    except Exception:  # a converter library choked on a malformed file: say so, never leak internals
        return None, "This file could not be read. It may be damaged or password protected."


def name_error(filename: str) -> str | None:
    """Why a file name cannot be imported, or None. Used by the file picker before anything is uploaded and by
    convert_upload, so the user sees one wording either way."""
    ext = PurePath(_clean_name(filename)).suffix.lower()
    if ext in LEGACY:
        return f"{ext} files are an old format. Save the file as {LEGACY[ext]} and upload it again."
    if ext not in SUPPORTED:
        return "Upload a Word (.docx), PowerPoint (.pptx), Excel (.xlsx), PDF, Markdown, HTML, CSV or text file."
    return None


def convert_upload(filename: str, data: bytes) -> Imported:
    name = _clean_name(filename)
    ext = PurePath(name).suffix.lower()
    problem = name_error(name)
    if problem:
        raise UnsupportedFile(problem)
    if not data:
        raise UnsupportedFile("The file is empty.")
    if len(data) > MAX_BYTES:
        raise UnsupportedFile(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")

    notes: list[str] = []
    if ext in (".docx", ".pptx", ".xlsx"):
        _check_office(data, ext)
        raw = {".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx}[ext](data, notes)
    elif ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise UnsupportedFile("This file is not a PDF.")
        raw = _pdf(data, notes)
    else:
        text = _decode(data)
        raw = {".md": _markdown, ".markdown": _markdown, ".html": _html, ".htm": _html, ".csv": _csv, ".txt": _text}[ext](
            text, notes)

    cleaned = _normalize(raw, notes, markdown=ext in (".md", ".markdown"))
    try:
        safe = sanitize_html(cleaned)
    except ValidationFailed:
        raise UnsupportedFile("The converted document is too large to open (1,000,000 characters at most).") from None
    plain = _plain_text(safe)
    if not plain.strip():
        if ext == ".pdf":
            raise UnsupportedFile("This PDF has no text layer (it looks scanned). Scanned pages cannot be imported.")
        raise UnsupportedFile("No content was found in this file.")
    return Imported(
        title=_title(name), html=safe, notes=tuple(notes), kind=ext.lstrip("."), words=len(plain.split()),
        headings=len(re.findall(r"<h[1-4][ >]", safe)), tables=safe.count("<table"))


# ---------------------------------------------------------------------------------------------------- safety checks

def _clean_name(filename: str) -> str:
    name = PurePath((filename or "").replace("\\", "/")).name
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    path = PurePath(name)
    stem, ext = path.stem[: 200 - len(path.suffix)], path.suffix  # shorten the stem, never the extension
    return (stem + ext) if stem else "document" + ext


def _title(name: str) -> str:
    stem = PurePath(name).stem
    stem = re.sub(r"[_]+", " ", stem).strip() or "Imported document"
    return stem[:MAX_TITLE]


def _decode(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise UnsupportedFile("This file looks binary, not text.")
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


REQUIRED_PART = {".docx": "word/document.xml", ".pptx": "ppt/presentation.xml", ".xlsx": "xl/workbook.xml"}


def _check_office(data: bytes, ext: str) -> None:
    """Archive guard plus a defusedxml parse of every XML part (refuses DTDs and entity expansion)."""
    from defusedxml import DefusedXmlException
    from defusedxml.ElementTree import fromstring

    if not data.startswith(b"PK\x03\x04"):
        raise UnsupportedFile(f"This is not a real {ext} file.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UnsupportedFile(f"This is not a real {ext} file.") from None
    infos = archive.infolist()
    if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise UnsupportedFile("This file is too large when unpacked, so it was refused.")
    if any(i.file_size > 1024 * 1024 and i.compress_size and i.file_size / i.compress_size > 200 for i in infos):
        raise UnsupportedFile("This file looks like a decompression bomb and was refused.")
    if REQUIRED_PART[ext] not in {i.filename for i in infos}:
        raise UnsupportedFile(f"This is not a real {ext} file.")
    for info in infos:
        if info.filename.endswith((".xml", ".rels")):
            if info.file_size > MAX_XML_PART:
                raise UnsupportedFile("This file contains an oversized part and was refused.")
            try:
                fromstring(archive.read(info))
            except DefusedXmlException:
                raise UnsupportedFile("This file contains unsafe XML and was refused.") from None
            except Exception:
                raise UnsupportedFile("This file is damaged.") from None


# ---------------------------------------------------------------------------------------------------- converters

def _esc(text: str) -> str:
    return html_lib.escape(text or "", quote=False)


def _table(rows: list[list[str]], notes: list[str], label: str = "") -> str:
    if not rows:
        return ""
    cut_rows, cut_cols = len(rows) > MAX_TABLE_ROWS, max(len(r) for r in rows) > MAX_TABLE_COLS
    if cut_rows or cut_cols:
        notes.append(f"{label + ': ' if label else ''}the table was cut to {MAX_TABLE_ROWS} rows and "
                     f"{MAX_TABLE_COLS} columns.")
    rows = [r[:MAX_TABLE_COLS] for r in rows[:MAX_TABLE_ROWS]]
    width = max(len(r) for r in rows)
    out = ["<table><tbody>"]
    for i, row in enumerate(rows):
        cell = "th" if i == 0 else "td"
        padded = row + [""] * (width - len(row))
        out.append("<tr>" + "".join(f"<{cell}><p>{_esc(c)}</p></{cell}>" for c in padded) + "</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def _docx(data: bytes, notes: list[str]) -> str:
    import mammoth

    skipped = 0

    def skip_image(_image):
        nonlocal skipped
        skipped += 1
        return {"src": "image-skipped"}  # the <img> is removed in _normalize

    style_map = "\n".join(["u => u", "strike => s", "p[style-name='Title'] => h1:fresh",
                           "p[style-name='Subtitle'] => p:fresh"])
    result = mammoth.convert_to_html(io.BytesIO(data), style_map=style_map,
                                     convert_image=mammoth.images.img_element(skip_image))
    return result.value


def _xlsx(data: bytes, notes: list[str]) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: list[str] = []
    sheets = [s for s in workbook.worksheets if s.sheet_state == "visible"]
    if len(sheets) > MAX_SHEETS:
        notes.append(f"Only the first {MAX_SHEETS} sheets were imported.")
    for sheet in sheets[:MAX_SHEETS]:
        rows: list[list[str]] = []
        for values in sheet.iter_rows(values_only=True, max_row=MAX_TABLE_ROWS + 1, max_col=MAX_TABLE_COLS + 1):
            cells = [_cell(v) for v in values]
            if any(cells):
                rows.append(cells)
        if rows:
            out.append(f"<h2>{_esc(sheet.title)}</h2>" + _table(rows, notes, sheet.title))
    if workbook.worksheets and any(s.sheet_state != "visible" for s in workbook.worksheets):
        notes.append("Hidden sheets were not imported.")
    notes.append("Excel values are imported as they were last saved; formulas are not kept.")
    return "".join(out)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ").replace(" 00:00:00", "")
    return str(value)


def _pptx(data: bytes, notes: list[str]) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(io.BytesIO(data))
    out: list[str] = []
    pictures = 0

    def walk(shapes, body: list[str], title_shape):
        nonlocal pictures
        for shape in shapes:
            if shape is title_shape:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, body, title_shape)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            elif getattr(shape, "has_table", False) and shape.has_table:
                body.append(_table([[c.text for c in row.cells] for row in shape.table.rows], notes))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                bullets = shape.is_placeholder or any(p.level for p in shape.text_frame.paragraphs)
                body.append(_pptx_text(shape.text_frame, bullets))

    slides = list(deck.slides)
    if len(slides) > MAX_SLIDES:
        notes.append(f"Only the first {MAX_SLIDES} slides were imported.")
    for number, slide in enumerate(slides[:MAX_SLIDES], 1):
        title_shape = slide.shapes.title
        title = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        body: list[str] = []
        walk(slide.shapes, body, title_shape)
        out.append(f"<h2>{_esc(title or f'Slide {number}')}</h2>" + "".join(body))
    if pictures:
        notes.append(f"{pictures} picture{'s' if pictures != 1 else ''} on slides {'were' if pictures != 1 else 'was'} "
                     "not imported.")
    return "".join(out)


def _pptx_text(frame, bullets: bool) -> str:
    items = [(p.level, "".join(r.text for r in p.runs).strip()) for p in frame.paragraphs]
    items = [(level, text) for level, text in items if text]
    if not bullets:
        return "".join(f"<p>{_esc(text)}</p>" for _, text in items)
    out: list[str] = []
    open_li: list[bool] = []  # one flag per open <ul>: is an <li> open at that depth?
    for level, text in items:
        target = min(level, 5) + 1
        while len(open_li) > target:
            if open_li.pop():
                out.append("</li>")
            out.append("</ul>")
        if len(open_li) == target:
            if open_li[-1]:
                out.append("</li>")
            open_li[-1] = False
        while len(open_li) < target:
            if open_li and not open_li[-1]:
                out.append("<li>")
                open_li[-1] = True
            out.append("<ul>")
            open_li.append(False)
        out.append(f"<li><p>{_esc(text)}</p>")
        open_li[-1] = True
    while open_li:
        if open_li.pop():
            out.append("</li>")
        out.append("</ul>")
    return "".join(out)


def _pdf(data: bytes, notes: list[str]) -> str:
    import pdfplumber

    lines_by_page: list[list[dict]] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = pdf.pages
        if len(pages) > MAX_PDF_PAGES:
            notes.append(f"Only the first {MAX_PDF_PAGES} pages were imported.")
        for page in pages[:MAX_PDF_PAGES]:
            words = page.extract_words(extra_attrs=["size"], use_text_flow=False, keep_blank_chars=False)
            lines: list[dict] = []
            for w in sorted(words, key=lambda w: (round(w["top"] / 3), w["x0"])):
                if lines and abs(lines[-1]["top"] - w["top"]) < 3:
                    lines[-1]["words"].append(w["text"])
                    lines[-1]["sizes"].append(w["size"])
                    lines[-1]["bottom"] = max(lines[-1]["bottom"], w["bottom"])
                else:
                    lines.append({"top": w["top"], "bottom": w["bottom"], "words": [w["text"]], "sizes": [w["size"]]})
            lines_by_page.append(lines)
    all_sizes = Counter(round(s) for page in lines_by_page for ln in page for s in ln["sizes"])
    if not all_sizes:
        return ""
    body_size = all_sizes.most_common(1)[0][0]
    out: list[str] = []
    paragraph: list[str] = []
    bullets: list[str] = []
    prev_bottom = None

    def flush():
        nonlocal paragraph, bullets
        if paragraph:
            out.append(f"<p>{_esc(' '.join(paragraph))}</p>")
            paragraph = []
        if bullets:
            out.append("<ul>" + "".join(f"<li><p>{_esc(b)}</p></li>" for b in bullets) + "</ul>")
            bullets = []

    for page in lines_by_page:
        prev_bottom = None
        for ln in page:
            text = " ".join(ln["words"]).strip()
            size = sum(ln["sizes"]) / len(ln["sizes"])
            if not text:
                continue
            gap = (ln["top"] - prev_bottom) if prev_bottom is not None else 0
            prev_bottom = ln["bottom"]
            if size >= body_size * 1.25 and len(text) <= 150:
                flush()
                level = 1 if size >= body_size * 1.7 else 2 if size >= body_size * 1.4 else 3
                out.append(f"<h{level}>{_esc(text)}</h{level}>")
            elif re.match(r"^[•●▪◦‣\-–*]\s+", text):
                if paragraph:
                    flush()
                bullets.append(re.sub(r"^[•●▪◦‣\-–*]\s+", "", text))
            else:
                if bullets:
                    flush()
                if paragraph and gap > body_size * 0.9:
                    flush()
                if paragraph and paragraph[-1].endswith("-") and not paragraph[-1].endswith(" -"):
                    paragraph[-1] = paragraph[-1][:-1] + text  # re-join a hyphenated word
                else:
                    paragraph.append(text)
        flush()
    notes.append("PDF tables are imported as plain text.")
    return "".join(out)


def _markdown(text: str, notes: list[str]) -> str:
    from markdown_it import MarkdownIt

    md = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
    return md.render(text)


def _html(text: str, notes: list[str]) -> str:
    try:
        doc = lxml_html.document_fromstring(text)
    except Exception:
        return f"<p>{_esc(text)}</p>"
    body = doc.find("body")
    root = body if body is not None else doc
    for bad in root.xpath("//script|//style|//iframe|//object|//embed|//form|//noscript|//template"):
        bad.drop_tree()
    return (_esc(root.text or "") if root.text and root.text.strip() else "") + "".join(
        lxml_html.tostring(child, encoding="unicode") for child in root)


def _csv(text: str, notes: list[str]) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [row for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]
    return _table(rows[: MAX_TABLE_ROWS + 1], notes) if rows else ""


def _text(text: str, notes: list[str]) -> str:
    blocks = [b for b in re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n")) if b.strip()]
    return "".join("<p>" + "<br>".join(_esc(line) for line in b.strip("\n").split("\n")) + "</p>" for b in blocks)


# ---------------------------------------------------------------------------------------------------- normalizing

def _normalize(raw: str, notes: list[str], *, markdown: bool = False) -> str:
    """Map whatever a converter produced onto the editor's schema, before the sanitizer enforces it."""
    if not raw.strip():
        return ""
    root = lxml_html.fragment_fromstring(raw, create_parent="div")

    images = 0
    for img in root.xpath(".//img|.//picture|.//svg|.//canvas|.//video|.//audio"):
        images += 1
        img.drop_tree()
    if images:
        notes.append(f"{images} image{'s' if images != 1 else ''} {'were' if images != 1 else 'was'} not imported "
                     "(images are not supported in documents).")

    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        tag = el.tag.lower()
        if tag in ("h5", "h6"):
            el.tag = "h4"
        elif tag == "b":
            el.tag = "strong"
        elif tag == "i":
            el.tag = "em"
        elif tag in ("strike", "del"):
            el.tag = "s"
    for el in root.xpath(".//font|.//center|.//ins|.//figure|.//figcaption|.//section|.//article|.//header|.//footer|.//nav|.//main"):
        el.drop_tag()

    if markdown:
        _markdown_blocks(root)

    for p in root.xpath("./p"):  # empty top-level paragraphs only add noise
        if not (p.text_content().strip() or len(p)):
            p.drop_tree()
    return "".join(
        (_esc(root.text) if root.text and root.text.strip() else "")
        + "".join(lxml_html.tostring(child, encoding="unicode") for child in root))


def _markdown_blocks(root) -> None:
    """Markdown extras: ```mermaid fences become diagram blocks, and GitHub '- [x]' items become checklists."""
    for pre in root.xpath(".//pre[code[contains(@class, 'language-mermaid')]]"):
        source = pre.find("code").text_content()
        div = lxml_html.Element("div")
        div.set("data-type", "mermaid")
        div.text = source.rstrip("\n")
        pre.getparent().replace(pre, div)
    for ul in root.xpath(".//ul"):
        items = ul.findall("li")
        marks = [re.match(r"^\s*\[( |x|X)\]\s+", _first_text(li)) for li in items]
        if items and all(marks):
            ul.set("data-type", "taskList")
            for li, mark in zip(items, marks):
                checked = mark.group(1).lower() == "x"
                _strip_marker(li, mark.end())
                li.set("data-type", "taskItem")
                li.set("data-checked", "true" if checked else "false")
                wrapper = lxml_html.Element("div")
                wrapper.text, li.text = li.text, None
                for child in list(li):
                    wrapper.append(child)
                li.append(wrapper)


def _first_text(li) -> str:
    if li.text and li.text.strip():
        return li.text
    first = li[0] if len(li) else None
    return (first.text or "") if first is not None and first.tag == "p" else ""


def _strip_marker(li, length: int) -> None:
    if li.text and li.text.strip():
        li.text = li.text[length:] if li.text.lstrip().startswith("[") else li.text
        li.text = re.sub(r"^\s*\[( |x|X)\]\s+", "", li.text)
    elif len(li) and li[0].tag == "p":
        li[0].text = re.sub(r"^\s*\[( |x|X)\]\s+", "", li[0].text or "")


def _plain_text(safe_html: str) -> str:
    if not safe_html.strip():
        return ""
    return lxml_html.fragment_fromstring(safe_html, create_parent="div").text_content()
