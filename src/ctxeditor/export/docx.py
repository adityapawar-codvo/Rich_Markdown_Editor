"""DOCX export: a python-docx walker over the sanitized editor HTML.

The result is editable in Word: real Heading 1-4 styles, real list numbering, real tables (merged cells, repeating
header rows), relationship-backed hyperlinks, and diagrams/charts as inline pictures (rendered by the browser and
passed in as PNG bytes). The metadata block, a DRAFT marker, the title header and the page-number footer are
generated here from the same DocMetadata the on-screen header uses.

Known limits: syntax-highlight colours come from Pygments, so they only approximate the on-screen theme; pictures are
scaled to at most 16 cm wide; fonts are Cambria/Calibri/Consolas, which Word substitutes if missing.
"""
from __future__ import annotations

import io
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass, field, replace
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.image.image import Image as DocxImage
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips
from lxml import html as lxml_html
from pygments.lexers import get_lexer_by_name
from pygments.styles import get_style_by_name
from pygments.util import ClassNotFound

from ..models import ExportPayload
from ..timefmt import format_local

BODY_FONT, UI_FONT, CODE_FONT = "Cambria", "Calibri", "Consolas"
TEXT_WIDTH = Cm(16)  # A4 (21 cm) minus 2 x 2.5 cm margins
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
HEADER_FILL = "F2F2F2"
CODE_FILL = "F5F5F5"
DEFAULT_HIGHLIGHT = "FFF59D"

BLOCK_TAGS = {"p", "h1", "h2", "h3", "h4", "ul", "ol", "li", "blockquote", "pre", "hr", "table", "div"}
HEADING_STYLES = {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3", "h4": "Heading 4"}
ALIGNMENTS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
              "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
URL_OK = re.compile(r"^(?:https?://|mailto:)", re.I)
BULLETS = ("•", "◦", "▪")
DECIMAL_FORMATS = (("decimal", "%{n}."), ("lowerLetter", "%{n}."), ("lowerRoman", "%{n}."))

# Schema order matters to Word, so raw elements are inserted before their successors.
PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                 "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                 "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                 "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
PPR_AFTER_SHD = PPR_AFTER_BDR[1:]
RPR_AFTER_SHD = ("w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
                 "w:specVanish", "w:oMath")
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")


@dataclass(frozen=True)
class _Fmt:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    sub: bool = False
    sup: bool = False
    code: bool = False
    color: str | None = None
    bg: str | None = None


@dataclass(frozen=True)
class _Ctx:
    container: object  # Document or _Cell
    para_style: str | None = None
    list_depth: int = 0
    in_cell: bool = False
    bold: bool = False
    max_width: object = TEXT_WIDTH
    factory: Callable[[], object] | None = None  # list items supply their own paragraph maker
    reserve: list = field(default_factory=list)  # a cell's initial empty paragraph, reused for the first block


# --------------------------------------------------------------------------- small helpers

def _hex(value: str) -> str | None:
    """#rgb, #rrggbb, #rrggbbaa, rgb() and rgba() as RRGGBB, or None."""
    value = value.strip().lower()
    if value.startswith("#"):
        digits = value[1:]
        if len(digits) in (3, 4):
            digits = "".join(ch * 2 for ch in digits[:3])
        if len(digits) in (6, 8) and re.fullmatch(r"[0-9a-f]+", digits):
            return digits[:6].upper()
        return None
    match = re.fullmatch(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*(?:,[^)]*)?\)", value)
    if match:
        return "".join(f"{min(int(g), 255):02X}" for g in match.groups())
    return None


def _style(el) -> dict[str, str]:
    out = {}
    for part in (el.get("style") or "").split(";"):
        name, _, raw = part.partition(":")
        if name.strip():
            out[name.strip().lower()] = raw.strip()
    return out


def _set_font(style, name: str) -> None:
    style.font.name = name
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):  # theme attributes override explicit names
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def _shd(fill: str):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    return el


def _shade_paragraph_pr(ppr, fill: str) -> None:
    ppr.insert_element_before(_shd(fill), *PPR_AFTER_SHD)


def _shade_run(run, fill: str) -> None:
    run._r.get_or_add_rPr().insert_element_before(_shd(fill), *RPR_AFTER_SHD)


def _shade_cell(cell, fill: str) -> None:
    cell._tc.get_or_add_tcPr().insert_element_before(_shd(fill), *TCPR_AFTER_SHD)


def _field(paragraph, instruction: str, font_size=None) -> None:
    """A complex field (PAGE, NUMPAGES) with a cached result of 1."""
    def piece(kind: str | None = None, text: str | None = None, instr: str | None = None):
        run = paragraph.add_run()
        if font_size:
            run.font.size = font_size
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        if instr:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = instr
            run._r.append(node)
        if text:
            run.text = text
    piece("begin")
    piece(instr=f" {instruction} ")
    piece("separate")
    piece(text="1")
    piece("end")


def _png_size(blob: bytes):
    image = DocxImage.from_blob(blob)
    return image.px_width, image.px_height


# --------------------------------------------------------------------------- builder

class _Builder:
    def __init__(self, doc: DocumentType, images: dict[int, bytes]):
        self.doc = doc
        self.images = images
        self.image_ordinal = 0
        self._configure()
        self._bullet_abstract = self._add_abstract(bullet=True, abstract_id=9001)
        self._decimal_abstract = self._add_abstract(bullet=False, abstract_id=9002)
        self._bullet_num: int | None = None

    # ---- document setup

    def _configure(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        _set_font(normal, BODY_FONT)
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for name, size, color, before in (("Heading 1", 20, "1F3864", 18), ("Heading 2", 16, "2F5496", 14),
                                          ("Heading 3", 13, "2F5496", 12), ("Heading 4", 11.5, "404040", 10)):
            style = styles[name]
            _set_font(style, UI_FONT)
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        title = styles["Title"]
        _set_font(title, UI_FONT)
        title.font.size = Pt(26)
        title.font.bold = True
        title.font.color.rgb = RGBColor.from_string("1F3864")
        quote = styles["Quote"]
        _set_font(quote, BODY_FONT)
        quote.font.italic = True
        quote.font.color.rgb = RGBColor.from_string("595959")
        quote.paragraph_format.left_indent = Cm(1)
        quote.paragraph_format.right_indent = Cm(1)
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.base_style = normal
        _set_font(code, CODE_FONT)
        code.font.size = Pt(9.5)
        code.paragraph_format.space_after = Pt(8)
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.left_indent = Cm(0.2)
        _shade_paragraph_pr(code.element.get_or_add_pPr(), CODE_FILL)
        for name in ("Header", "Footer"):
            styles[name].paragraph_format.tab_stops.clear_all()
        styles["Header"].paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
        for style_name in ("Header", "Footer"):
            _set_font(styles[style_name], UI_FONT)
            styles[style_name].font.size = Pt(9)
            styles[style_name].font.color.rgb = RGBColor.from_string("595959")

        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.5))

    # ---- numbering

    def _numbering(self):
        return self.doc.part.numbering_part.element

    def _add_abstract(self, *, bullet: bool, abstract_id: int) -> int:
        levels = []
        for lvl in range(9):
            if bullet:
                fmt, text = "bullet", BULLETS[lvl % 3]
                font = '<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:hint="default"/>'
            else:
                fmt, template = DECIMAL_FORMATS[lvl % 3]
                text, font = template.format(n=lvl + 1), ""
            levels.append(
                f'<w:lvl w:ilvl="{lvl}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                f'<w:pPr><w:ind w:left="{720 + 360 * lvl}" w:hanging="360"/></w:pPr>'
                f"<w:rPr>{font}</w:rPr></w:lvl>")
        node = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(levels)}</w:abstractNum>')
        numbering = self._numbering()
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(node)  # abstractNum elements must precede num elements
        else:
            numbering.append(node)
        return abstract_id

    def _new_num(self, abstract_id: int, *, restart_level: int | None = None, start: int = 1) -> int:
        numbering = self._numbering()
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        num_id = max(ids, default=0) + 1
        override = ""
        if restart_level is not None:
            override = (f'<w:lvlOverride w:ilvl="{restart_level}"><w:startOverride w:val="{start}"/>'
                        "</w:lvlOverride>")
        numbering.append(parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override}</w:num>'))
        return num_id

    # ---- top of document

    def build(self, payload: ExportPayload, tz_name: str | None) -> None:
        meta = payload.metadata
        props = self.doc.core_properties
        props.title, props.author, props.last_modified_by = meta.title, meta.last_edited_by, meta.last_edited_by
        props.subject = meta.doc_type
        props.created = props.modified = datetime.now(UTC).replace(tzinfo=None)
        self._header_footer(payload)
        if payload.is_draft:
            banner = self.doc.add_paragraph()
            run = banner.add_run("DRAFT")
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = UI_FONT
            run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
            banner.add_run("  This version is not published.").font.name = UI_FONT
        self.doc.add_paragraph(meta.title, style="Title")
        self._metadata_table(payload, tz_name)
        self.doc.add_paragraph()
        html = (payload.html or "").strip()
        if html:
            root = lxml_html.fragment_fromstring(html, create_parent="div")
            self._flow(_Ctx(container=self.doc), root)
        body = self.doc.element.body
        children = [c for c in body if c.tag != qn("w:sectPr")]
        if children and children[-1].tag == qn("w:tbl"):
            self.doc.add_paragraph()  # Word requires a paragraph after a final table

    def _header_footer(self, payload: ExportPayload) -> None:
        section = self.doc.sections[0]
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0]
        paragraph.text = ""
        paragraph.add_run(payload.metadata.title)
        if payload.is_draft:
            paragraph.add_run("\t")
            run = paragraph.add_run("DRAFT")
            run.bold = True
            run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run("Page ")
        _field(paragraph, "PAGE")
        paragraph.add_run(" of ")
        _field(paragraph, "NUMPAGES")

    def _metadata_table(self, payload: ExportPayload, tz_name: str | None) -> None:
        meta = payload.metadata
        rows = (("Title", meta.title), ("Type", meta.doc_type), ("Version", str(meta.version_number)),
                ("Status", meta.status), ("Created", format_local(meta.created_at, tz_name)),
                ("Created by", meta.created_by), ("Last edited", format_local(meta.last_edited_at, tz_name)),
                ("Last edited by", meta.last_edited_by),
                ("Last published", format_local(meta.last_published_at, tz_name)),
                ("Published by", meta.last_published_by or "-"))
        table = self.doc.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        table.autofit = False
        for row, (label, value) in zip(table.rows, rows):
            for cell, width in zip(row.cells, (Cm(4), Cm(12))):
                cell.width = width
            first, second = row.cells
            _shade_cell(first, HEADER_FILL)
            for cell, text, bold in ((first, label, True), (second, value, False)):
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                run = paragraph.add_run(text)
                run.bold = bold
                run.font.name = UI_FONT
                run.font.size = Pt(10)

    # ---- block structure

    @staticmethod
    def _items(el) -> Iterator:
        """el's text and children in order, each child followed by its tail."""
        if el.text:
            yield el.text
        for child in el:
            yield child
            if child.tail:
                yield child.tail

    def _flow(self, ctx: _Ctx, el) -> None:
        """Walk el's content: block children become blocks, runs of inline content become paragraphs."""
        pending: list = []

        def flush() -> None:
            if pending and any(isinstance(i, str) and i.strip() or not isinstance(i, str) for i in pending):
                self._inline_paragraph(ctx, pending[:])
            pending.clear()

        for item in self._items(el):
            if isinstance(item, str):
                pending.append(item)
            elif not isinstance(item.tag, str):  # comment or processing instruction
                continue
            elif item.tag in BLOCK_TAGS:
                flush()
                self._block(ctx, item)
            else:
                pending.append(item)
        flush()

    def _block(self, ctx: _Ctx, el) -> None:
        tag = el.tag
        if tag == "p":
            self._paragraph(ctx, el)
        elif tag in HEADING_STYLES:
            paragraph = self._make_para(replace(ctx, para_style=HEADING_STYLES[tag]))
            self._align(paragraph, el)
            self._fill(paragraph, ctx, list(self._items(el)))
        elif tag in ("ul", "ol"):
            self._list(ctx, el)
        elif tag == "blockquote":
            self._flow(replace(ctx, para_style="Quote", factory=None), el)
        elif tag == "pre":
            self._code_block(ctx, el)
        elif tag == "hr":
            self._rule(ctx)
        elif tag == "table":
            self._table(ctx, el)
        elif tag == "div" and el.get("data-type") in ("mermaid", "chart"):
            self._picture(ctx, el)
        else:  # li outside a list, div, anything unknown: walk as a container
            self._flow(ctx, el)

    def _make_para(self, ctx: _Ctx):
        if ctx.reserve:
            paragraph = ctx.reserve.pop()
            if ctx.para_style:
                paragraph.style = ctx.para_style
        else:
            paragraph = ctx.container.add_paragraph(style=ctx.para_style)
        if ctx.in_cell:
            paragraph.paragraph_format.space_after = Pt(3)
        return paragraph

    def _new_para(self, ctx: _Ctx):
        return ctx.factory() if ctx.factory else self._make_para(ctx)

    @staticmethod
    def _align(paragraph, el) -> None:
        align = _style(el).get("text-align", "").lower()
        if align in ALIGNMENTS:
            paragraph.alignment = ALIGNMENTS[align]

    def _paragraph(self, ctx: _Ctx, el) -> None:
        paragraph = self._new_para(ctx)
        self._align(paragraph, el)
        self._fill(paragraph, ctx, list(self._items(el)))

    def _inline_paragraph(self, ctx: _Ctx, items: list) -> None:
        self._fill(self._new_para(ctx), ctx, items)

    # ---- inline content

    def _collect(self, item, fmt: _Fmt, link: str | None, out: list) -> None:
        if isinstance(item, str):
            out.append(("t", item, fmt, link))
            return
        tag = item.tag
        if not isinstance(tag, str):
            return
        if tag == "br":
            out.append(("br",))
            return
        style = _style(item)
        if tag in ("strong", "b"):
            fmt = replace(fmt, bold=True)
        elif tag in ("em", "i"):
            fmt = replace(fmt, italic=True)
        elif tag == "u":
            fmt = replace(fmt, underline=True)
        elif tag in ("s", "strike", "del"):
            fmt = replace(fmt, strike=True)
        elif tag == "sub":
            fmt = replace(fmt, sub=True)
        elif tag == "sup":
            fmt = replace(fmt, sup=True)
        elif tag == "code":
            fmt = replace(fmt, code=True)
        elif tag == "mark":
            fmt = replace(fmt, bg=_hex(style.get("background-color", "")) or DEFAULT_HIGHLIGHT)
        elif tag == "a":
            href = (item.get("href") or "").strip()
            if URL_OK.match(href):
                link = href
        if tag == "span" or tag == "mark":
            color = _hex(style.get("color", ""))
            if color:
                fmt = replace(fmt, color=color)
            if tag == "span":
                bg = _hex(style.get("background-color", ""))
                if bg:
                    fmt = replace(fmt, bg=bg)
        for child in self._items(item):
            self._collect(child, fmt, link, out)

    def _fill(self, paragraph, ctx: _Ctx, items: list) -> None:
        segments: list = []
        for item in items:
            self._collect(item, _Fmt(bold=ctx.bold), None, segments)
        # HTML whitespace rules: collapse runs, drop leading and trailing blanks of the paragraph.
        previous_space = True
        for index, seg in enumerate(segments):
            if seg[0] == "br":
                previous_space = True
                continue
            text = re.sub(r"\s+", " ", seg[1])
            if previous_space:
                text = text.lstrip()
            if text:
                previous_space = text.endswith(" ")
            segments[index] = ("t", text, seg[2], seg[3])
        for index in range(len(segments) - 1, -1, -1):  # trailing blanks of the paragraph
            if segments[index][0] == "t" and segments[index][1]:
                segments[index] = ("t", segments[index][1].rstrip(), *segments[index][2:])
                break
        hyperlink, hyperlink_url = None, None
        for seg in segments:
            if seg[0] == "br":
                paragraph.add_run().add_break()
                hyperlink, hyperlink_url = None, None
                continue
            _, text, fmt, link = seg
            if not text:
                continue
            if link and link == hyperlink_url and hyperlink is not None:
                container = hyperlink
            elif link:
                hyperlink = self._hyperlink(paragraph, link)
                hyperlink_url, container = link, hyperlink
            else:
                hyperlink, hyperlink_url, container = None, None, None
            run = paragraph.add_run(text)
            self._format_run(run, fmt, linked=bool(link))
            if container is not None:
                container.append(run._r)

    @staticmethod
    def _hyperlink(paragraph, url: str):
        rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        node = OxmlElement("w:hyperlink")
        node.set(qn("r:id"), rel_id)
        node.set(qn("w:history"), "1")
        paragraph._p.append(node)
        return node

    @staticmethod
    def _format_run(run, fmt: _Fmt, *, linked: bool = False) -> None:
        font = run.font
        if fmt.bold:
            font.bold = True
        if fmt.italic:
            font.italic = True
        if fmt.underline or linked:
            font.underline = True
        if fmt.strike:
            font.strike = True
        if fmt.sub:
            font.subscript = True
        if fmt.sup:
            font.superscript = True
        if fmt.code:
            font.name = CODE_FONT
            font.size = Pt(10)
        if linked:
            font.color.rgb = LINK_COLOR
        if fmt.color:
            font.color.rgb = RGBColor.from_string(fmt.color)
        fill = fmt.bg or (CODE_FILL if fmt.code else None)
        if fill:
            _shade_run(run, fill)

    # ---- lists

    def _list(self, ctx: _Ctx, el) -> None:
        depth = ctx.list_depth
        task = el.get("data-type") == "taskList"
        ordered = el.tag == "ol"
        num_id = None
        if ordered:
            start = el.get("start") or "1"
            num_id = self._new_num(self._decimal_abstract, restart_level=depth,
                                   start=int(start) if start.isdigit() else 1)
        elif not task:
            if self._bullet_num is None:
                self._bullet_num = self._new_num(self._bullet_abstract)
            num_id = self._bullet_num
        for li in el.findall("li"):
            checked = li.get("data-checked") == "true"
            self._list_item(ctx, li, depth, ordered, num_id, task, checked)

    def _list_item(self, ctx: _Ctx, li, depth: int, ordered: bool, num_id: int | None, task: bool,
                   checked: bool) -> None:
        state = {"used": False}
        base = replace(ctx, para_style=None, factory=None)
        left = Twips(720 + 360 * depth)

        def factory():
            paragraph = self._make_para(base)
            if not state["used"]:
                state["used"] = True
                paragraph.paragraph_format.left_indent = left
                paragraph.paragraph_format.first_line_indent = Twips(-360)
                paragraph.paragraph_format.space_after = Pt(3)
                if task:
                    paragraph.add_run("☑ " if checked else "☐ ")
                else:
                    kind = "Number" if ordered else "Bullet"
                    level = "" if depth == 0 else f" {min(depth + 1, 3)}"
                    paragraph.style = f"List {kind}{level}"
                    numpr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                    numpr.get_or_add_ilvl().val = min(depth, 8)
                    numpr.get_or_add_numId().val = num_id
                    paragraph.paragraph_format.left_indent = left
                    paragraph.paragraph_format.first_line_indent = Twips(-360)
            else:
                paragraph.paragraph_format.left_indent = left
            return paragraph

        self._flow(replace(ctx, factory=factory, list_depth=depth + 1), li)
        if not state["used"]:
            factory()

    # ---- code, rule, pictures

    def _code_block(self, ctx: _Ctx, el) -> None:
        code = el.find("code")
        source = (code if code is not None else el).text_content()
        source = source[:-1] if source.endswith("\n") else source
        language = None
        for cls in ((code.get("class") if code is not None else None) or "").split():
            if cls.startswith("language-"):
                language = cls[len("language-"):]
        paragraph = self._make_para(replace(ctx, para_style="Code Block"))
        tokens = self._tokens(source, language)
        for text, color, bold, italic in tokens:
            for number, line in enumerate(text.split("\n")):
                if number:
                    paragraph.add_run().add_break()
                if line:
                    run = paragraph.add_run(line)
                    if color:
                        run.font.color.rgb = RGBColor.from_string(color)
                    if bold:
                        run.font.bold = True
                    if italic:
                        run.font.italic = True

    @staticmethod
    def _tokens(source: str, language: str | None) -> list[tuple[str, str | None, bool, bool]]:
        if language:
            try:
                lexer = get_lexer_by_name(language)
            except ClassNotFound:
                lexer = None
            if lexer is not None:
                style = get_style_by_name("default")
                out = []
                for ttype, value in lexer.get_tokens(source):
                    spec = style.style_for_token(ttype)
                    out.append((value, spec.get("color"), bool(spec.get("bold")), bool(spec.get("italic"))))
                while out and out[-1][0].endswith("\n"):  # the lexer appends a final newline
                    value = out[-1][0].rstrip("\n")
                    if value:
                        out[-1] = (value, *out[-1][1:])
                        break
                    out.pop()
                return out
        return [(source, None, False, False)]

    def _rule(self, ctx: _Ctx) -> None:
        paragraph = self._make_para(replace(ctx, para_style=None))
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "A0A0A0")):
            bottom.set(qn(f"w:{key}"), value)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().insert_element_before(border, "w:shd", *PPR_AFTER_SHD)

    def _picture(self, ctx: _Ctx, el) -> None:
        ordinal, self.image_ordinal = self.image_ordinal, self.image_ordinal + 1
        kind = el.get("data-type")
        source = el.text_content() if kind == "mermaid" else (el.get("data-chart") or "")
        title = "Diagram" if kind == "mermaid" else "Chart"
        blob = self.images.get(ordinal)
        if blob:
            try:
                px_width, _ = _png_size(blob)
                width = min(ctx.max_width, Inches(px_width / 96))
                paragraph = self._make_para(replace(ctx, para_style=None))
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                run.add_picture(io.BytesIO(blob), width=width)
                doc_pr = run._r.xpath(".//wp:docPr")[0]
                doc_pr.set("descr", source)
                doc_pr.set("title", title)
                return
            except Exception:  # an unreadable image falls back to the placeholder
                pass
        note = self._make_para(replace(ctx, para_style=None))
        label = title if kind == "mermaid" else self._chart_title(source)
        run = note.add_run(f"[{title} image unavailable" + (f": {label}]" if kind == "chart" and label else "]"))
        run.italic = True
        run.font.color.rgb = RGBColor.from_string("595959")
        if kind == "mermaid" and source.strip():
            paragraph = self._make_para(replace(ctx, para_style="Code Block"))
            for number, line in enumerate(source.strip("\n").split("\n")):
                if number:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)

    @staticmethod
    def _chart_title(raw: str) -> str:
        match = re.search(r'"title"\s*:\s*"((?:[^"\\]|\\.)*)"', raw)
        return match.group(1) if match else ""

    # ---- tables

    def _table(self, ctx: _Ctx, el) -> None:
        trs = el.xpath("./thead/tr | ./tbody/tr | ./tfoot/tr | ./tr")
        if not trs:
            return
        occupied: set[tuple[int, int]] = set()
        placements = []
        for r, tr in enumerate(trs):
            c = 0
            for cell_el in tr.xpath("./th | ./td"):
                while (r, c) in occupied:
                    c += 1
                colspan = self._span(cell_el.get("colspan"), 50)
                rowspan = self._span(cell_el.get("rowspan"), len(trs) - r)
                while colspan > 1 and any((r, c + k) in occupied for k in range(colspan)):
                    colspan -= 1  # never overlap an earlier rowspan
                for rr in range(r, r + rowspan):
                    for cc in range(c, c + colspan):
                        occupied.add((rr, cc))
                placements.append((r, c, colspan, rowspan, cell_el))
                c += colspan
        if not placements:
            return
        ncols = max(c + cs for _, c, cs, _, _ in placements)
        header_rows = {r for r, tr in enumerate(trs)
                       if tr.getparent().tag == "thead" or (tr.xpath("./th") and not tr.xpath("./td"))}
        table = ctx.container.add_table(rows=len(trs), cols=ncols)
        table.style = "Table Grid"
        table.autofit = False
        col_width = int(ctx.max_width / ncols)
        for row in table.rows:
            for cell in row.cells:
                cell.width = col_width
        for r, c, colspan, rowspan, _ in placements:
            if colspan > 1 or rowspan > 1:
                table.cell(r, c).merge(table.cell(r + rowspan - 1, c + colspan - 1))
        for r in header_rows:
            trpr = table.rows[r]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:tblHeader"))
        for r, c, colspan, rowspan, cell_el in placements:
            cell = table.cell(r, c)
            is_header = r in header_rows or cell_el.tag == "th"
            if is_header:
                _shade_cell(cell, HEADER_FILL)
            initial = cell.paragraphs[0]
            width = int(col_width * colspan) - Cm(0.4)
            inner = replace(ctx, container=cell, in_cell=True, bold=is_header, list_depth=0, factory=None,
                            para_style=None, reserve=[initial], max_width=max(width, Cm(2)))
            self._flow(inner, cell_el)
            if inner.reserve and len(cell._tc.findall(qn("w:p"))) > 1:
                cell._tc.remove(initial._p)  # drop the unused empty first paragraph

    @staticmethod
    def _span(raw: str | None, limit: int) -> int:
        return max(1, min(int(raw), limit)) if raw and raw.isdigit() else 1


def build_docx(payload: ExportPayload, images: dict[int, bytes], tz_name: str | None = None) -> bytes:
    """A .docx for the payload. images maps the 0-based ordinal of each mermaid/chart block (both kinds counted
    together, in document order) to PNG bytes; a missing image becomes a placeholder, never an error."""
    doc = Document()
    _Builder(doc, images or {}).build(payload, tz_name)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
